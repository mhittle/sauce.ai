"""FDA device submissions (510(k) + De Novo) per condition, and their public
summary documents.

- **Sync** (per condition): pull every openFDA ``device/510k`` record whose
  device name matches the condition's terms (the same proxy that drives the
  directory ranking), store the full public record, and count 510(k)s and
  De Novos separately (De Novo numbers start with ``DEN``; openFDA serves
  both from this endpoint).
- **Documents** (background worker): download each submission's public
  summary PDF (510(k) Summary / De Novo decision summary) from
  accessdata.fda.gov, extract its text, list the predicate devices it names,
  and match it against the catalog — dataset aliases, descriptor-paper
  titles and DOIs — so datasets and papers "used in a 510(k)" are flagged.
- **LLM summary** (on demand by default): when a user opens a submission,
  Claude turns the PDF into a structured review (intended use, training /
  test data, performance, predicates, datasets and references named).
"""
from __future__ import annotations

import io
import json
import logging
import re
import threading
import time
from typing import Any

import requests

from . import fda
from .config import Settings
from .literature import find_dois
from .store import Store, normalize_condition

log = logging.getLogger(__name__)
logging.getLogger("pypdf").setLevel(logging.ERROR)  # font-encoding chatter

FDA_DOCS = "https://www.accessdata.fda.gov/cdrh_docs"
FDA_DB = "https://www.accessdata.fda.gov/scripts/cdrh/cfdocs/cfpmn"
_UA = {"User-Agent": "Mozilla/5.0 (compatible; sauce.ai-datasets/0.1; +https://sauce.ai)"}
_SUB_RE = re.compile(r"\b(K\d{6}|DEN\d{6})\b")
PAGE = 1000
MAX_PDF_BYTES = 40_000_000
MAX_PAGES = 120
MAX_TEXT = 400_000
LLM_TEXT = 150_000


# ------------------------------------------------------------------ software / AI
# Software: FDA generic names / device names / regulations that describe
# software functions (CAD, image processing, analysis algorithms, apps).
_SOFTWARE_RE = re.compile(
    r"software|algorithm|computer[- ]?(aided|assisted)|\bcad[ex]?\b|image processing"
    r"|image analysis|artificial intelligence|machine learning|deep learning|\bai\b"
    r"|mobile medical app|\bapp\b|analysis system|decision support|triage and notification",
    re.I)
_SOFTWARE_REGULATIONS = {
    "892.2050",  # medical image management and processing system
    "892.2060", "892.2070", "892.2080", "892.2090",  # CADx, CADe, triage, CADe/x
    "892.2100",  # radiological acquisition / optimization guidance
    "870.2785",  # software for optimizing cardiac therapy
    "882.1491",  # pediatric autism spectrum disorder diagnosis aid
    "870.2345",  # ECG software for over-the-counter use
    "870.2380",  # cardiovascular machine learning-based notification software
}
# AI: phrases in the summary PDF that indicate a learned model.
_AI_TEXT_RE = re.compile(
    r"artificial intelligence|machine[- ]learning|deep[- ]learning|neural network"
    r"|convolutional|trained (?:model|algorithm|network)|\bAI[- ](?:based|enabled|algorithm)",
    re.I)


# Displays/monitors/workstations share regulation 892.2050 and the "image
# processing system" generic name with software; they're hardware.
_HARDWARE_RE = re.compile(r"\b(monitor|display|lcd|oled|workstation|viewing station)s?\b", re.I)
_STRONG_SOFTWARE_RE = re.compile(
    r"software|algorithm|artificial intelligence|machine learning|deep learning|\bai\b"
    r"|computer[- ]?(aided|assisted)|\bcad[ex]?\b|\bapp\b", re.I)


def is_software_record(d: dict) -> bool:
    name = d.get("device_name") or ""
    generic = d.get("generic_name") or ""
    if _STRONG_SOFTWARE_RE.search(name):
        return True
    if _HARDWARE_RE.search(name) or _HARDWARE_RE.search(generic):
        return False
    return (bool(_SOFTWARE_RE.search(f"{generic} {name}"))
            or d.get("regulation_number") in _SOFTWARE_REGULATIONS)


def text_says_ai(text: str) -> bool:
    return len(_AI_TEXT_RE.findall(text[:200_000])) >= 2  # one passing mention isn't enough


# ------------------------------------------------------------------ records
def submission_type(k: str) -> str:
    k = k.upper()
    return "denovo" if k.startswith("DEN") else "pma" if k.startswith("P") else "510k"


def normalize_record(r: dict) -> dict:
    ofda = r.get("openfda") or {}
    first = lambda v: (v[0] if isinstance(v, list) and v else v) or None  # noqa: E731
    address = ", ".join(x for x in (r.get("address_1"), r.get("address_2"), r.get("city"),
                                    r.get("state"), r.get("zip_code") or r.get("postal_code"),
                                    r.get("country_code")) if x)
    k = r["k_number"]
    return {
        "k_number": k, "submission_type": submission_type(k),
        "device_name": r.get("device_name"), "applicant": r.get("applicant"),
        "contact": re.sub(r"\s+", " ", r.get("contact") or "").strip() or None,
        "address": address or None, "country_code": r.get("country_code"),
        "decision_date": r.get("decision_date"), "decision_code": r.get("decision_code"),
        "decision_description": r.get("decision_description"),
        "date_received": r.get("date_received"), "product_code": r.get("product_code"),
        "generic_name": first(ofda.get("device_name")),
        "device_class": first(ofda.get("device_class")),
        "regulation_number": first(ofda.get("regulation_number")),
        "medical_specialty": first(ofda.get("medical_specialty_description"))
                             or r.get("advisory_committee_description"),
        "advisory_committee": r.get("advisory_committee_description"),
        "clearance_type": r.get("clearance_type"),
        "statement_or_summary": r.get("statement_or_summary"),
        "third_party_flag": r.get("third_party_flag"),
        "expedited_review_flag": r.get("expedited_review_flag"),
        "is_software": None,  # set below
        "raw": json.dumps({k2: v for k2, v in r.items() if k2 != "openfda"}
                          | {"openfda": {k2: v for k2, v in ofda.items()
                                         if k2 not in ("registration_number", "fei_number")}}),
    }


def _classified(d: dict) -> dict:
    d["is_software"] = is_software_record(d)
    return d


def fda_database_url(k: str) -> str:
    t = submission_type(k)
    if t == "pma":
        return f"https://www.accessdata.fda.gov/scripts/cdrh/cfdocs/cfpma/pma.cfm?id={k}"
    return f"{FDA_DB}/{'denovo.cfm' if t == 'denovo' else 'pmn.cfm'}?ID={k}"


def summary_pdf_urls(k: str) -> list[str]:
    """Candidate locations of the public summary PDF. 510(k)s live under
    pdf<yy>/ (unpadded year, 2002+) or pdf/ (older); De Novos under reviews/."""
    k = k.upper()
    if k.startswith("DEN"):
        return [f"{FDA_DOCS}/reviews/{k}.pdf"]
    if k.startswith("P"):  # PMA: Summary of Safety and Effectiveness Data ("B" file)
        yy = int(k[1:3]) if k[1:3].isdigit() else 0
        return [f"{FDA_DOCS}/pdf{yy}/{k}B.pdf"] if 2 <= yy < 70 else [f"{FDA_DOCS}/pdf/{k}B.pdf"]
    try:
        yy = int(k[1:3])
    except ValueError:
        return [f"{FDA_DOCS}/pdf/{k}.pdf"]
    urls = [f"{FDA_DOCS}/pdf{yy}/{k}.pdf"] if 2 <= yy < 70 else []
    return urls + [f"{FDA_DOCS}/pdf/{k}.pdf"]


# ------------------------------------------------------------------ sync
def sync_condition(store: Store, name: str, settings: Settings,
                   session: Any = None) -> dict | None:
    """Fetch all matching submissions for one condition. Returns counts, or
    None on an openFDA failure (existing data is kept)."""
    cond = store.get_condition(name)
    if not cond:
        return None
    terms = [t for t in [cond["name"], *cond["fda_terms"]] if len(t) >= 4]
    expr = fda.build_search(terms)
    http = session or requests
    key = f"&api_key={settings.openfda_api_key}" if settings.openfda_api_key else ""
    records: list[dict] = []
    total, skip = 0, 0
    while expr:
        url = (f"{fda.API}?search={expr}&limit={PAGE}&skip={skip}"
               f"&sort=decision_date:desc{key}")
        try:
            resp = http.get(url, headers=_UA, timeout=30)
        except requests.RequestException as exc:
            log.warning("openFDA sync %s failed: %s", name, exc)
            return None
        if resp.status_code == 404:
            break  # no matches
        if resp.status_code != 200:
            log.warning("openFDA %s for %s", resp.status_code, name)
            return None
        body = resp.json()
        total = int(body["meta"]["results"]["total"])
        batch = body.get("results", [])
        records += batch
        skip += len(batch)
        if not batch or skip >= total or skip >= settings.fda_max_devices_per_condition:
            break
    devices = [_classified(normalize_record(r)) for r in records if r.get("k_number")]
    store.upsert_devices(devices)
    n_denovo = sum(d["submission_type"] == "denovo" for d in devices)
    capped = total > len(devices)
    # When capped (rare: generic terms), De Novos are counted among the stored
    # most-recent records and everything else is attributed to 510(k).
    n_510k = total - n_denovo
    store.set_condition_devices(cond["name"], [d["k_number"] for d in devices],
                                n_510k, n_denovo, capped)
    return {"n_510k": n_510k, "n_denovo": n_denovo, "stored": len(devices), "total": total}


def fetch_one(store: Store, k: str, settings: Settings, session: Any = None) -> bool:
    """Fetch a single submission by number (e.g. a predicate we haven't
    stored because it matched no condition). True if found."""
    if not re.fullmatch(r"(K\d{6}|DEN\d{6})", k):
        return False
    http = session or requests
    key = f"&api_key={settings.openfda_api_key}" if settings.openfda_api_key else ""
    try:
        resp = http.get(f"{fda.API}?search=k_number:{k}&limit=1{key}", headers=_UA, timeout=30)
    except requests.RequestException:
        return False
    if resp.status_code != 200:
        return False
    results = resp.json().get("results") or []
    if not results:
        return False
    store.upsert_devices([_classified(normalize_record(results[0]))])
    return True


# ------------------------------------------------------------------ FDA AI list
AI_LIST_KEY = "fda_ai_list_synced_at"
_NUM_RE = re.compile(r"\b(K\d{6}|DEN\d{6}|P\d{6}(?:/S\d+)?|H\d{6})\b")


def parse_ai_list(data: bytes) -> list[dict]:
    """FDA "AI-Enabled Medical Devices" spreadsheet → rows. Tolerates column
    order changes (matched by header text) and HYPERLINK formulas."""
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data), read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows)]

    def col(*names):
        for i, h in enumerate(header):
            if any(n in h for n in names):
                return i
        return None
    ci = {"number": col("submission"), "date": col("date"), "device": col("device"),
          "company": col("company", "applicant"), "panel": col("panel"),
          "product_code": col("product code")}
    if ci["number"] is None:
        raise ValueError("AI list: no submission-number column")
    out = []
    for r in rows:
        m = _NUM_RE.search(str(r[ci["number"]] or ""))
        if not m:
            continue
        get = lambda key: (str(r[ci[key]]).strip() if ci[key] is not None  # noqa: E731
                           and r[ci[key]] is not None else None)
        date = get("date")
        if date and re.fullmatch(r"\d{2}/\d{2}/\d{4}", date):
            mm, dd, yy = date.split("/")
            date = f"{yy}-{mm}-{dd}"
        elif date and " " in date:  # datetime cell
            date = date.split(" ")[0]
        out.append({"k_number": m.group(1).split("/")[0], "decision_date": date,
                    "device_name": get("device"), "applicant": get("company"),
                    "advisory_committee": get("panel"), "product_code": get("product_code")})
    return out


def sync_ai_list(store: Store, settings: Settings, session: Any = None) -> int | None:
    """Download FDA's AI-enabled device list, add every listed submission to
    the catalog (stub rows, then full openFDA records in batches), and flag
    them AI. Returns the number listed, or None on failure."""
    http = session or requests
    try:
        resp = http.get(settings.fda_ai_list_url, headers=_UA, timeout=60)
        if resp.status_code != 200:
            raise ValueError(f"HTTP {resp.status_code}")
        entries = parse_ai_list(resp.content)
    except Exception as exc:
        log.warning("FDA AI list sync failed: %s", exc)
        return None
    stubs = []
    for e in entries:
        k = e["k_number"]
        stype = ("denovo" if k.startswith("DEN") else "510k" if k.startswith("K") else "pma")
        stubs.append({**e, "submission_type": stype, "is_software": True})
    store.insert_device_stubs(stubs)  # never overwrites a full openFDA record
    store.mark_ai([s["k_number"] for s in stubs], "fda_list")
    # Enrich K/DEN numbers we don't have full records for (50 per request).
    missing = store.missing_devices([s["k_number"] for s in stubs
                                     if s["submission_type"] != "pma"])
    key = f"&api_key={settings.openfda_api_key}" if settings.openfda_api_key else ""
    for i in range(0, len(missing), 50):
        batch = missing[i:i + 50]
        try:
            r = http.get(f"{fda.API}?search=k_number:({'+'.join(batch)})&limit=100{key}",
                         headers=_UA, timeout=30)
        except requests.RequestException:
            break
        if r.status_code == 200:
            store.upsert_devices([_classified(normalize_record(x))
                                  for x in r.json().get("results", [])])
        elif r.status_code == 429:
            break  # daily limit; the rest fill in on the next sync
    store.kv_set(AI_LIST_KEY, _now_iso())
    return len(stubs)


def _now_iso() -> str:
    from .store import now_iso
    return now_iso()


# ------------------------------------------------------------------ AI list → conditions
MAP_BATCH = 40
MAP_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["devices"],
    "properties": {"devices": {"type": "array", "items": {
        "type": "object", "additionalProperties": False,
        "required": ["k_number", "conditions"],
        "properties": {"k_number": {"type": "string"},
                       "conditions": {"type": "array", "items": {"type": "string"}}}}}},
}
MAP_SYSTEM = """You assign FDA-cleared AI-enabled medical devices to the clinical conditions they detect, diagnose, triage, monitor or treat, so a catalog can count devices per condition.

Rules:
- Condition names are lowercase and unabbreviated ("intracranial hemorrhage", not "ICH"; "atrial fibrillation", not "AFib"). Use an existing condition name from the list whenever it fits; create a new one only when none does.
- Use the disease or clinical finding level ("breast cancer", "pulmonary embolism", "diabetic retinopathy", "lung nodule"), not the modality or body part.
- 0-4 conditions per device. Return [] for condition-agnostic tools (organ/anatomy segmentation, image reconstruction or denoising, workflow, PACS, dose tracking, generic measurement) and when the name and codes don't tell you — never guess.
- Return every k_number you were given, exactly as given."""


def map_devices_to_conditions(client, settings: Settings, devices: list[dict],
                              known_conditions: list[str]) -> dict[str, list[str]]:
    from .crawler import _create
    lines = [json.dumps({k: d.get(k) for k in (
        "k_number", "device_name", "applicant", "generic_name", "product_code",
        "advisory_committee", "medical_specialty", "regulation_number") if d.get(k)})
        for d in devices]
    resp = _create(
        client, settings, system=MAP_SYSTEM,
        output_config={"effort": "low",
                       "format": {"type": "json_schema", "schema": MAP_SCHEMA}},
        messages=[{"role": "user", "content":
                   "Existing conditions:\n" + ", ".join(known_conditions[:400])
                   + "\n\nDevices (one JSON object per line):\n" + "\n".join(lines)}])
    if resp.stop_reason == "refusal":
        raise RuntimeError("mapping request declined")
    out = json.loads(next(b.text for b in resp.content if b.type == "text"))
    given = {d["k_number"] for d in devices}
    mapping = {}
    for row in out.get("devices", []):
        if row.get("k_number") in given:
            mapping[row["k_number"]] = [c.strip().lower() for c in row.get("conditions", [])
                                        if isinstance(c, str) and 3 <= len(c.strip()) <= 80][:4]
    return mapping


# ------------------------------------------------------------------ documents
def fetch_summary_pdf(k: str, session: Any = None) -> tuple[bytes | None, str | None, str | None]:
    """(pdf bytes, url, error). Tries each candidate location."""
    http = session or requests
    last = "no summary PDF found"
    for url in summary_pdf_urls(k):
        try:
            resp = http.get(url, headers=_UA, timeout=60, stream=True)
        except requests.RequestException as exc:
            last = f"network: {exc}"[:200]
            continue
        if resp.status_code == 404:
            continue
        if resp.status_code != 200:
            last = f"HTTP {resp.status_code}"
            continue
        ctype = resp.headers.get("Content-Type", "")
        if "pdf" not in ctype.lower():
            last = f"not a PDF ({ctype})"
            continue
        buf = bytearray()
        for chunk in resp.iter_content(1 << 16):
            buf.extend(chunk)
            if len(buf) > MAX_PDF_BYTES:
                return None, url, "PDF too large"
        return bytes(buf), url, None
    return None, None, last


def pdf_text(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages[:MAX_PAGES]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # one bad page shouldn't lose the document
            parts.append("")
    return "\n".join(parts)[:MAX_TEXT], len(reader.pages)


def _norm(s: str) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", s.lower()))


def _snippet(text: str, start: int, end: int) -> str:
    return re.sub(r"\s+", " ", text[max(0, start - 160):end + 160]).strip()


def find_links(store: Store, text: str, extra_names: list[str] = ()) -> list[dict]:
    """Catalog datasets/papers referenced in a submission summary."""
    targets = store.link_targets()
    links: dict[tuple, dict] = {}
    for ds_id, alias in targets["aliases"]:
        pat = re.compile(r"(?<![A-Za-z0-9])" + re.escape(alias) + r"(?![A-Za-z0-9])", re.I)
        m = pat.search(text)
        if m:
            links.setdefault((ds_id, ""), {"dataset_id": ds_id, "via": "alias",
                                          "matched": alias,
                                          "snippet": _snippet(text, m.start(), m.end())})
        elif any(alias.lower() == n.lower() for n in extra_names):
            links.setdefault((ds_id, ""), {"dataset_id": ds_id, "via": "alias",
                                          "matched": alias, "snippet": None})
    norm_text = _norm(text)
    for ds_id, art_id, title in targets["titles"]:
        nt = _norm(title)[:120]
        if len(nt) >= 30 and nt in norm_text:
            links[(ds_id, art_id)] = {"dataset_id": ds_id, "article_id": art_id,
                                      "via": "title", "matched": title[:200], "snippet": None}
    for doi in find_dois(text):
        art = store.find_article_id(doi=doi)
        if art:
            for ds_id in store.datasets_for_article(art):
                pos = text.lower().find(doi)
                links[(ds_id, art)] = {"dataset_id": ds_id, "article_id": art, "via": "doi",
                                       "matched": doi,
                                       "snippet": _snippet(text, pos, pos + len(doi))
                                       if pos >= 0 else None}
    return list(links.values())


EXTRACT_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["summary", "intended_use", "technology", "is_ai_ml", "training_data",
                 "test_data", "reference_standard", "performance", "predicate_devices",
                 "datasets_named", "references", "limitations"],
    "properties": {
        "summary": {"type": "string"},
        "intended_use": {"type": "string"},
        "technology": {"type": "string"},
        "is_ai_ml": {"type": "boolean"},
        "training_data": {"type": "string"},
        "test_data": {
            "type": "object", "additionalProperties": False,
            "required": ["description", "n_cases", "n_sites", "countries", "design"],
            "properties": {k: {"type": "string"} for k in
                           ("description", "n_cases", "n_sites", "countries", "design")},
        },
        "reference_standard": {"type": "string"},
        "performance": {
            "type": "array",
            "items": {"type": "object", "additionalProperties": False,
                      "required": ["metric", "value", "ci", "population"],
                      "properties": {k: {"type": "string"} for k in
                                     ("metric", "value", "ci", "population")}},
        },
        "predicate_devices": {"type": "array", "items": {"type": "string"}},
        "datasets_named": {"type": "array", "items": {"type": "string"}},
        "references": {"type": "array", "items": {"type": "string"}},
        "limitations": {"type": "string"},
    },
}

EXTRACT_SYSTEM = """You review public FDA device submission summaries (510(k) \
Summaries and De Novo decision summaries) for people building software-as-a-medical-\
device. Extract only what the document states; use an empty string or empty list when \
it says nothing. Never infer numbers.

- summary: 2-3 plain sentences on what the device does and how it was validated.
- technology: how it works (e.g. "CNN on chest radiographs"); is_ai_ml true if it uses \
machine learning / AI.
- training_data / test_data: sources, sizes, sites, countries, retrospective vs \
prospective, as stated.
- performance: each reported metric with its value, confidence interval and the \
population/subgroup it applies to.
- predicate_devices: K/DEN numbers and names of predicates.
- datasets_named: any named datasets or public datasets (e.g. "NIH ChestX-ray14").
- references: bibliographic references cited, one string each."""


def llm_extract(client, settings: Settings, device: dict, text: str) -> dict:
    from .crawler import _create
    header = {k: device.get(k) for k in ("k_number", "device_name", "applicant",
                                         "decision_date", "product_code", "generic_name")}
    resp = _create(
        client, settings, system=EXTRACT_SYSTEM,
        output_config={"effort": "low",
                       "format": {"type": "json_schema", "schema": EXTRACT_SCHEMA}},
        messages=[{"role": "user", "content":
                   f"Submission: {json.dumps(header)}\n\nDocument text:\n{text[:LLM_TEXT]}"}])
    if resp.stop_reason == "refusal":
        raise RuntimeError("extraction request declined")
    return json.loads(next(b.text for b in resp.content if b.type == "text"))


# ------------------------------------------------------------------ worker
def _older_than(iso: str, days: int) -> bool:
    from datetime import datetime, timedelta, timezone
    try:
        return datetime.fromisoformat(iso) < datetime.now(timezone.utc) - timedelta(days=days)
    except ValueError:
        return True


def _days_ago_iso(days: float) -> str:
    from datetime import datetime, timedelta, timezone
    return (datetime.now(timezone.utc) - timedelta(days=days)).replace(microsecond=0).isoformat()


class DeviceWorker:
    """Background loop: sync one stale condition, or read one summary PDF, or
    run one requested LLM extraction, per step."""

    def __init__(self, store: Store, settings: Settings, client_factory=None,
                 session: Any = None):
        self.store = store
        self.settings = settings
        self.client_factory = client_factory
        self.session = session
        self._client = None
        self._mapper = None
        self._stop = threading.Event()
        self._wake = threading.Event()
        self._thread: threading.Thread | None = None
        self.last_error: str | None = None

    def start(self) -> None:
        if self._thread and self._thread.is_alive():
            return
        self._stop.clear()
        self._thread = threading.Thread(target=self._loop, name="devices", daemon=True)
        self._thread.start()

    def stop(self) -> None:
        self._stop.set()
        self._wake.set()

    def poke(self) -> None:
        """A user asked about a submission: skip the idle wait."""
        self._wake.set()

    def running(self) -> bool:
        return bool(self._thread and self._thread.is_alive())

    def _loop(self) -> None:
        while not self._stop.is_set():
            try:
                worked = self.step()
            except Exception as exc:
                log.exception("device step failed")
                self.last_error = f"{type(exc).__name__}: {str(exc)[:300]}"
                worked = True
            self._wake.wait(self.settings.device_interval_sec if worked
                            else self.settings.lit_idle_sec)
            self._wake.clear()

    def _llm(self):
        if self.settings.device_llm == "off":
            return None
        if self._client is None and self.client_factory is not None:
            self._client = self.client_factory(self.settings)
        return self._client

    def step(self) -> bool:
        # User requests beat background work.
        job = self.store.next_device_doc(want_extraction=self._llm() is not None)
        if job and self._is_request(job[0]):
            return self.run_job(*job)
        if not self.store.kv_get("software_classified_v2"):
            self.store.reclassify_software(is_software_record)
            self.store.kv_set("software_classified_v2", _now_iso())
        last = self.store.kv_get(AI_LIST_KEY)
        if self.settings.fda_ai_list_url and (not last or _older_than(last, days=7)):
            sync_ai_list(self.store, self.settings, self.session)
            if not self.store.kv_get(AI_LIST_KEY):  # failed: don't retry every step
                self.store.kv_set(AI_LIST_KEY, _days_ago_iso(6))
            return True
        if self.settings.device_mapping and self.client_factory is not None:
            batch = self.store.devices_to_map(MAP_BATCH)
            if batch and not self._map_paused():
                self.map_batch(batch)
                return True
        cond = self.store.condition_due_for_device_sync(self.settings.lit_refresh_days)
        if cond:
            if sync_condition(self.store, cond, self.settings, self.session) is None:
                # Mark as attempted so one failing condition can't starve the rest.
                c = self.store.get_condition(cond)
                self.store.set_condition_devices(
                    cond, self._condition_ks(cond), c.get("fda_510k_count") or 0,
                    c.get("fda_denovo_count") or 0)
            return True
        if job:
            return self.run_job(*job)
        return False

    def _map_paused(self) -> bool:
        until = self.store.kv_get("device_mapping_paused_until")
        return bool(until) and not _older_than(until, days=0)

    def map_batch(self, batch: list[dict]) -> None:
        if self._mapper is None:
            self._mapper = self.client_factory(self.settings)
        try:
            mapping = map_devices_to_conditions(self._mapper, self.settings, batch,
                                                self.store.condition_names())
        except Exception as exc:
            # API trouble: back off an hour instead of retrying every step.
            log.warning("device→condition mapping failed: %s", exc)
            self.last_error = f"mapping: {type(exc).__name__}: {str(exc)[:200]}"
            self.store.kv_set("device_mapping_paused_until", _days_ago_iso(-1 / 24))
            return
        # A device the model left out counts as condition-agnostic ([]), so a
        # malformed answer can't keep the same batch coming back every step.
        self.store.save_device_conditions({d["k_number"]: mapping.get(d["k_number"], [])
                                           for d in batch})

    def _is_request(self, k: str) -> bool:
        with self.store._conn() as c:
            row = c.execute("SELECT requested FROM device_docs WHERE k_number = ?",
                            (k,)).fetchone()
        return bool(row and row["requested"])

    def _condition_ks(self, cond: str) -> list[str]:
        with self.store._conn() as c:
            return [r["k_number"] for r in c.execute(
                "SELECT k_number FROM condition_devices WHERE condition = ?",
                (normalize_condition(cond),))]

    def run_job(self, k: str, job: str) -> bool:
        if job == "fetch":
            self.read_document(k)
            wanted = self._is_request(k) or self.settings.device_llm == "all"
            if wanted and self._llm() is not None:
                self.extract(k)
        elif job == "extract":
            self.extract(k)
        return True

    def read_document(self, k: str) -> None:
        data, url, err = fetch_summary_pdf(k, self.session)
        if data is None:
            status = "none" if err == "no summary PDF found" else "failed"
            self.store.save_device_doc(k, status, pdf_url=url, error=err)
            if self._is_request(k):
                self.store.save_device_extraction(k, None, error=err)
            return
        try:
            text, pages = pdf_text(data)
        except Exception as exc:
            self.store.save_device_doc(k, "failed", pdf_url=url,
                                       error=f"PDF parse: {type(exc).__name__}")
            return
        preds = sorted({m.group(1) for m in _SUB_RE.finditer(text)} - {k})
        self.store.save_device_doc(k, "done", pdf_url=url, pages=pages, text=text,
                                   predicates=preds)
        if text_says_ai(text):
            self.store.mark_ai([k], "summary_text")
        self.store.replace_device_links(k, find_links(self.store, text))

    def extract(self, k: str) -> None:
        client = self._llm()
        text = self.store.device_text(k)
        if not client or not text:
            self.store.save_device_extraction(k, None)
            return
        try:
            out = llm_extract(client, self.settings, self.store.get_device(k), text)
        except Exception as exc:
            log.warning("extraction failed for %s: %s", k, exc)
            self.store.save_device_extraction(k, None, error=f"summary failed: {type(exc).__name__}")
            return
        self.store.save_device_extraction(k, out)
        if out.get("is_ai_ml"):
            self.store.mark_ai([k], "summary_review")
        # Names the model found can confirm catalog links the regex missed.
        links = find_links(self.store, text, extra_names=out.get("datasets_named", []))
        self.store.replace_device_links(k, links)

    def run_for(self, seconds: float) -> int:
        end, steps = time.monotonic() + seconds, 0
        while time.monotonic() < end and self.step():
            steps += 1
            time.sleep(self.settings.device_interval_sec)
        return steps
